import io
import openpyxl
from pathlib import Path
from app.config import settings
from app.models.vendor import Vendor

TAX_ID_COLUMNS = ["J", "K", "L", "M", "N", "O", "P", "Q", "R", "S", "T", "U", "V"]

def export_vendor_to_excel(vendor: Vendor) -> io.BytesIO:
    """
    Populate vendor data into the official Excel template.
    Returns BytesIO containing the populated workbook.
    """
    wb = openpyxl.load_workbook(str(settings.TEMPLATE_PATH))
    ws = wb.active

    # 1. Header & Requestor Info
    if vendor.req_date:
        ws["C4"] = f"…… {vendor.req_date} ……  Email Address: …… {vendor.req_email or ''} ……"
    if vendor.req_department:
        ws["R4"] = vendor.req_department
    if vendor.req_bu:
        ws["AB4"] = vendor.req_bu
    if vendor.req_name:
        ws["C5"] = f"…… {vendor.req_name} ……"
    if vendor.req_phone:
        ws["O5"] = vendor.req_phone
    if vendor.req_ext:
        ws["AC5"] = vendor.req_ext

    # 2. Status & Code
    if vendor.vendor_code:
        ws["S8"] = vendor.vendor_code

    # 4. Name and Address
    if vendor.vendor_name_en:
        ws["B15"] = f"  ภาษาอังกฤษ : {vendor.vendor_name_en}"
    if vendor.vendor_name_th:
        ws["B16"] = f"  ภาษาไทย : {vendor.vendor_name_th}"

    ws["B17"] = (
        f"  ที่อยู่ :  เลขที่ {vendor.address_no or '-'}  หมู่ที่ {vendor.address_moo or '-'}  "
        f"อาคาร {vendor.address_building or '-'}  ชั้น {vendor.address_floor or '-'}  หมู่บ้าน {vendor.address_village or '-'}"
    )
    ws["B18"] = (
        f"  ตรอก/ซอย {vendor.address_soi or '-'}  ถนน {vendor.address_road or '-'}  "
        f"ตำบล/แขวง {vendor.address_subdistrict or '-'}  อำเภอ/เขต {vendor.address_district or '-'}"
    )
    ws["B19"] = (
        f"  จังหวัด {vendor.address_province or '-'}  รหัสไปรษณีย์ {vendor.address_postcode or '-'}  "
        f"โทรศัพท์ {vendor.vendor_phone or '-'}  โทรสาร {vendor.vendor_fax or '-'}"
    )

    if vendor.email_etax:
        ws["C21"] = f"…… {vendor.email_etax} ……"

    # 5. Tax ID (13 boxes)
    digits = [c for c in (vendor.tax_id or "") if c.isdigit()]
    for i, col in enumerate(TAX_ID_COLUMNS):
        if i < len(digits):
            ws[f"{col}22"] = int(digits[i])

    if vendor.branch_no:
        ws["Z22"] = vendor.branch_no

    # 6. Contact Person
    if vendor.contact_name:
        ws["B24"] = (
            f"  ชื่อผู้สามารถติดต่อ : {vendor.contact_name}   โทรศัพท์ : {vendor.contact_phone or '-'}   "
            f"มือถือ : {vendor.contact_mobile or '-'}   E-Mail : {vendor.contact_email or '-'}"
        )

    # 12. Bank Details
    if vendor.bank_account_name_th:
        ws["B34"] = f"     ชื่อบัญชี (ไทย) : {vendor.bank_account_name_th}"
    if vendor.bank_account_name_en:
        ws["B35"] = f"     ชื่อบัญชี (อังกฤษ) : {vendor.bank_account_name_en}"
    if vendor.bank_name:
        ws["B36"] = f"     ธนาคาร : {vendor.bank_name}"
    if vendor.bank_branch:
        ws["B37"] = f"     สาขา : {vendor.bank_branch}"
    if vendor.bank_account_no:
        ws["Q37"] = vendor.bank_account_no

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
